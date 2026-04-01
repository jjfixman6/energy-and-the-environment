"""
SMR (NuScale) Presentation Generator
Generates a professional 16:9 PowerPoint with dark theme and speaker notes.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# --- Color Palette ---
BG_COLOR = RGBColor(0x1B, 0x1F, 0x3B)       # Dark navy
TITLE_COLOR = RGBColor(0xFF, 0xFF, 0xFF)     # White
BULLET_COLOR = RGBColor(0xD0, 0xD4, 0xE4)   # Light grey-blue
ACCENT_COLOR = RGBColor(0x4E, 0xC9, 0xB0)   # Teal accent
ACCENT2_COLOR = RGBColor(0xF0, 0xA5, 0x30)  # Amber accent
SUBTLE_COLOR = RGBColor(0x8A, 0x8E, 0xA8)   # Muted text
SLIDE_NUM_COLOR = RGBColor(0x60, 0x64, 0x80) # Dim slide numbers

# Standard 16:9 dimensions in EMUs
SLIDE_WIDTH = Emu(12192000)
SLIDE_HEIGHT = Emu(6858000)

# --- Slide Data ---
SLIDES = [
    {
        "title": "The Baseload Problem",
        "bullets": [
            "Global electricity demand: 2x by 2050",
            "Renewables: intermittent, no 24/7 guarantee",
            "Conventional nuclear: $20-30B, 15+ years",
        ],
        "footer": "Can we get nuclear reliability without the cost?",
        "notes": (
            "By 2050, global electricity demand is projected to double. At the same time, "
            "we need to cut carbon emissions by roughly fifty percent this decade. Renewables "
            "are essential — but solar and wind are intermittent. They cannot guarantee power "
            "at 2am in January. Conventional nuclear can — but it costs twenty to thirty "
            "billion dollars per plant and takes over fifteen years to build. So the question "
            "is: can we get nuclear's reliability without its cost and complexity?"
        ),
    },
    {
        "title": "A New Manufacturing Model",
        "bullets": [
            "< 300 MW, factory-built modules",
            "Standardized components, shipped to site",
            "Incremental deployment as demand grows",
        ],
        "footer": "Not a smaller reactor — a different construction paradigm",
        "notes": (
            "That's the premise behind Small Modular Reactors. An SMR isn't just a smaller "
            "reactor — it's a fundamentally different manufacturing model. A conventional "
            "plant exceeds 1,000 megawatts and is built entirely on-site — custom-engineered, "
            "slow, and expensive. An SMR is under 300 megawatts, factory-fabricated under "
            "controlled conditions, then shipped and assembled on-site. That distinction is "
            "critical. Factory production means standardized components, tighter quality "
            "control, and learning-curve cost reductions — the same logic that transformed "
            "aerospace and shipbuilding. And crucially, you don't commit twenty billion "
            "dollars upfront. You deploy modules incrementally, matching capacity to demand."
        ),
    },
    {
        "title": "How It Works: Fission to Grid",
        "bullets": [
            "U-235 fission → E = mc² energy release",
            "Heat → pressurized water → steam generator",
            "Rankine cycle: 30-35% thermal efficiency",
        ],
        "footer": "Not a design flaw — just physics",
        "notes": (
            "The underlying physics is identical to any nuclear plant. Uranium-235 undergoes "
            "fission — the nucleus splits, and because the fission products have slightly "
            "less mass than the original atom, that mass difference is converted to energy, "
            "exactly as E equals mc squared predicts. That energy heats pressurized water — "
            "held above 300 degrees Celsius by maintaining high pressure — which transfers "
            "thermal energy to a secondary loop through a steam generator. The steam drives "
            "a turbine connected to a generator. Thermodynamic efficiency sits at 30 to 35 "
            "percent — that's the Rankine cycle ceiling for this temperature range. Not a "
            "design flaw — just physics."
        ),
    },
    {
        "title": "NuScale: Passive Safety",
        "bullets": [
            "Natural circulation: no pumps needed",
            "Gravity-driven cooling — no power required",
            "Decay heat removed by thermodynamics alone",
        ],
        "footer": "The safety case is physics, not procedure",
        "notes": (
            "Now, here is where NuScale's design makes a genuinely important engineering "
            "contribution. Every nuclear reactor — even after shutdown — continues producing "
            "decay heat from fission product decay. In a conventional plant, you remove that "
            "heat with electrically powered pumps. But pumps need power, and power can fail. "
            "Fukushima showed exactly what happens when it does. NuScale eliminates that "
            "single point of failure entirely. Its cooling system relies on natural "
            "circulation: heated coolant is less dense, so it rises; cooler fluid descends "
            "to replace it. This is a convection loop driven by thermodynamics alone — no "
            "pumps, no external power, no operator action required. The safety case is "
            "physics, not procedure."
        ),
    },
    {
        "title": "Why NuScale Leads",
        "bullets": [
            "First & only NRC-certified SMR design",
            "Integral PWR: core + steam gen + pressurizer",
            "77 MWe per module → 462 MW (6 modules)",
        ],
        "footer": "Not a concept — a certified, deployable design",
        "notes": (
            "NuScale is the most deployment-ready SMR design in the world. It is the first "
            "— and currently only — SMR to receive design certification from the U.S. "
            "Nuclear Regulatory Commission, arguably the most rigorous nuclear regulator on "
            "earth. Its integral pressurized water reactor design consolidates the reactor "
            "core, steam generator, and pressurizer into a single sealed pressure vessel — "
            "eliminating external piping and the failure points that come with it. Each "
            "module produces 77 megawatts electric. Six modules give you a 462 megawatt plant."
        ),
    },
    {
        "title": "The Economic Reality",
        "bullets": [
            "Capital cost: $6,000-$10,000/kW installed",
            "LCOE: $65-90/MWh vs. solar at $30-40",
            "Idaho project cancelled — FOAK risk is real",
        ],
        "footer": "Loss of economies of scale is the core structural problem",
        "notes": (
            "Now the hard part — and I want to be honest about this. Capital costs for SMRs "
            "are currently estimated at six to ten thousand dollars per kilowatt installed. "
            "That is not cheaper than conventional nuclear — it's comparable or worse on a "
            "per-kilowatt basis. Projected levelized cost of energy sits at 65 to 90 dollars "
            "per megawatt-hour. That beats natural gas when you factor in carbon pricing, but "
            "it cannot compete with utility-scale solar at 30 to 40 dollars. The core problem "
            "is loss of economies of scale — larger plants spread fixed costs across more "
            "megawatts. SMRs don't get that advantage yet. And the cancellation of NuScale's "
            "flagship Idaho project confirmed that first-of-a-kind cost risk is real, not "
            "theoretical."
        ),
    },
    {
        "title": "Barriers to Deployment",
        "bullets": [
            "Regulatory: licensing not built for modular",
            "Supply chain: factory model unproven at scale",
            "Market: solar/wind costs still falling",
        ],
        "footer": "Nth-of-a-kind economics require the first plants to get built",
        "notes": (
            "Three structural barriers remain. First, regulatory: even with NRC "
            "certification, project-level licensing is slow and expensive — the regulatory "
            "framework was not designed for modular, repeated deployment. Second, supply "
            "chain: the factory manufacturing model that makes SMRs attractive doesn't exist "
            "at scale yet, so the cost advantages remain largely theoretical. Third, market "
            "competition: solar and wind costs keep falling. SMRs need to achieve "
            "nth-of-a-kind economics — where the tenth or twentieth unit benefits from "
            "accumulated manufacturing learning — to close that gap. That requires the first "
            "plants to actually get built."
        ),
    },
    {
        "title": "Where SMRs Fit",
        "bullets": [
            "Not a silver bullet — a specific tool",
            "Firm, dispatchable, low-carbon baseload",
            "Engineering proven; barrier is economic & solvable",
        ],
        "footer": "The grid needs what renewables alone cannot provide",
        "notes": (
            "So where does that leave us? Small Modular Reactors are not a silver bullet. "
            "They will not out-compete solar on cost, and they will not solve the storage "
            "problem that renewables face. But what they offer is something very specific: "
            "firm, dispatchable, low-carbon baseload power — the kind of generation the grid "
            "needs when the wind stops blowing and the sun goes down. The engineering is "
            "proven. The passive safety is real. The remaining barrier is economic — and it "
            "is solvable, if the first projects get financed and built."
        ),
    },
]


def set_slide_bg(slide, color):
    """Set solid background color for a slide."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, shape_type, left, top, width, height, fill_color=None, line_color=None):
    """Add a shape to a slide."""
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.background()
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_textbox(slide, left, top, width, height, text, font_size, color,
                bold=False, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    """Add a text box to a slide."""
    txbox = slide.shapes.add_textbox(left, top, width, height)
    tf = txbox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txbox


def get_blank_layout(prs):
    """Find blank layout safely."""
    for layout in prs.slide_layouts:
        if layout.name == "Blank":
            return layout
    # Fallback: use the last layout or the one with fewest placeholders
    return min(prs.slide_layouts, key=lambda l: len(l.placeholders))


def build_slide(prs, slide_data, slide_num, total_slides):
    """Build a single content slide."""
    slide_layout = get_blank_layout(prs)
    slide = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide, BG_COLOR)

    # Accent bar at top
    add_shape(slide, MSO_SHAPE.RECTANGLE,
              Inches(0), Inches(0), SLIDE_WIDTH, Inches(0.06),
              fill_color=ACCENT_COLOR)

    # Title
    add_textbox(slide, Inches(0.8), Inches(0.5), Inches(11), Inches(0.9),
                slide_data["title"], 36, TITLE_COLOR, bold=True)

    # Thin separator line under title
    add_shape(slide, MSO_SHAPE.RECTANGLE,
              Inches(0.8), Inches(1.35), Inches(2.5), Inches(0.04),
              fill_color=ACCENT_COLOR)

    # Bullets
    bullet_top = Inches(1.8)
    for i, bullet_text in enumerate(slide_data["bullets"]):
        # Bullet marker
        add_textbox(slide, Inches(0.8), bullet_top + Inches(i * 0.65),
                    Inches(0.4), Inches(0.5),
                    "\u25B8", 18, ACCENT_COLOR)
        # Bullet text
        add_textbox(slide, Inches(1.2), bullet_top + Inches(i * 0.65),
                    Inches(5.5), Inches(0.55),
                    bullet_text, 20, BULLET_COLOR, font_name="Calibri Light")

    # Visual placeholder area (right side)
    placeholder = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
                            Inches(7.2), Inches(1.5), Inches(5.3), Inches(4.2),
                            line_color=RGBColor(0x3A, 0x3E, 0x5C))
    # Placeholder label
    add_textbox(slide, Inches(7.5), Inches(3.2), Inches(4.8), Inches(0.8),
                "[Visual: see slides.md for specs]", 14, SUBTLE_COLOR,
                alignment=PP_ALIGN.CENTER, font_name="Calibri Light")

    # Footer / key takeaway
    if slide_data.get("footer"):
        add_textbox(slide, Inches(0.8), Inches(6.2), Inches(11), Inches(0.6),
                    slide_data["footer"], 16, ACCENT_COLOR, bold=False,
                    font_name="Calibri Light", alignment=PP_ALIGN.LEFT)

    # Slide number
    add_textbox(slide, Inches(12.0), Inches(6.8), Inches(1), Inches(0.4),
                f"{slide_num}/{total_slides}", 11, SLIDE_NUM_COLOR,
                alignment=PP_ALIGN.RIGHT, font_name="Calibri Light")

    # Speaker notes
    notes_slide = slide.notes_slide
    notes_tf = notes_slide.notes_text_frame
    notes_tf.text = slide_data["notes"]

    return slide


def build_title_slide(prs):
    """Build the opening title slide (before S1)."""
    slide_layout = get_blank_layout(prs)
    slide = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide, BG_COLOR)

    # Accent bar
    add_shape(slide, MSO_SHAPE.RECTANGLE,
              Inches(0), Inches(0), SLIDE_WIDTH, Inches(0.06),
              fill_color=ACCENT_COLOR)

    # Main title
    add_textbox(slide, Inches(1.5), Inches(2.0), Inches(10), Inches(1.5),
                "Small Modular Reactors", 48, TITLE_COLOR, bold=True,
                alignment=PP_ALIGN.CENTER)

    # Subtitle
    add_textbox(slide, Inches(1.5), Inches(3.5), Inches(10), Inches(1.0),
                "Engineering, Economics, and the Case for NuScale", 24,
                ACCENT_COLOR, alignment=PP_ALIGN.CENTER,
                font_name="Calibri Light")

    # Separator
    add_shape(slide, MSO_SHAPE.RECTANGLE,
              Inches(5.5), Inches(4.6), Inches(2.3), Inches(0.04),
              fill_color=ACCENT_COLOR)

    # Presentation info
    add_textbox(slide, Inches(1.5), Inches(5.0), Inches(10), Inches(0.6),
                "Energy and the Environment", 18, SUBTLE_COLOR,
                alignment=PP_ALIGN.CENTER, font_name="Calibri Light")

    # Speaker notes
    notes_slide = slide.notes_slide
    notes_tf = notes_slide.notes_text_frame
    notes_tf.text = "Title slide — no script. Advance immediately after introduction."

    return slide


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    # Title slide
    build_title_slide(prs)

    # Content slides
    total = len(SLIDES)
    for i, slide_data in enumerate(SLIDES, 1):
        build_slide(prs, slide_data, i, total)

    output_path = "SMR_NuScale_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved to: {output_path}")
    print(f"Total slides: {total + 1} (1 title + {total} content)")


if __name__ == "__main__":
    main()
